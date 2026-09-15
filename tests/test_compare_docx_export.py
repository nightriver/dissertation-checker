from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Twips
from lxml import etree

from compare.docx_export import (
    COLUMN_WIDTHS,
    DOCUMENT_TITLE,
    EMPTY_TEXT,
    MISSING_PAGE_MARKER,
    build_comparison_docx,
    page_marker,
)
from compare.matcher import compare_tokens
from compare.normalize import tokenize_lines
from compare.presentation import LINE_BREAK, fragment_pieces
from table_highlighter.processor import process_document
from table_highlighter.types import HighlightOptions


def _lines(text, page=None):
    return [{"line": text, "page": page}]


def _pair():
    base = [f"термін{i}" for i in range(30)]
    changed = base[:15] + ["заміна"] + base[16:]
    lines_a, lines_b = _lines(" ".join(base), 3), _lines(" ".join(changed), 7)
    tokens_a, tokens_b = tokenize_lines(lines_a), tokenize_lines(lines_b)
    segments = compare_tokens(tokens_a, tokens_b).segments
    return segments, lines_a, tokens_a, lines_b, tokens_b


def _open(data):
    return Document(BytesIO(data))


def _build(**kwargs):
    segments, lines_a, tokens_a, lines_b, tokens_b = _pair()
    return build_comparison_docx(
        segments, lines_a, tokens_a, lines_b, tokens_b, name_a="a", name_b="b", **kwargs
    )


def _highlight(run):
    element = run._r.find(qn("w:rPr")).find(qn("w:highlight"))
    return None if element is None else element.get(qn("w:val"))


def _xml(element):
    return etree.tostring(element) if element is not None else None


def _highlighter_reference(font_name, font_size):
    """Та сама пара фрагментів, пропущена через справжній table-highlight."""
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("термін0 термін1", "термін0 заміна")):
        for index, width in enumerate(COLUMN_WIDTHS):
            table.rows[0].cells[index].width = Twips(width)
        cell.paragraphs[0].text = "С. 3"
        cell.add_paragraph(text)
    buffer = BytesIO()
    document.save(buffer)
    options = HighlightOptions(font_name=font_name, font_size=font_size)
    return _open(process_document(buffer.getvalue(), options).document_bytes).tables[0]


def test_docx_has_title_summary_and_one_row_per_finding():
    segments, lines_a, tokens_a, lines_b, tokens_b = _pair()
    document = _open(build_comparison_docx(
        segments, lines_a, tokens_a, lines_b, tokens_b,
        name_a="дисертація.pdf", name_b="джерело.docx", summary=["Знайдені фрагменти: 1"],
    ))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert DOCUMENT_TITLE in text
    assert "дисертація.pdf" in text and "джерело.docx" in text
    assert "Знайдені фрагменти: 1" in text

    table = document.tables[0]
    # Ні шапки, ні рядків показників: лише пари фрагментів.
    assert len(table.rows) == len(segments)
    left, right = table.rows[0].cells
    assert left._tc is not right._tc
    assert [paragraph.text for paragraph in left.paragraphs][0] == "С. 3"
    assert [paragraph.text for paragraph in right.paragraphs][0] == "С. 7"
    assert "аркуш" not in left.text and "змінений" not in left.text
    assert "термін0" in left.text and "заміна" not in left.text
    assert "заміна" in right.text


def test_docx_page_geometry_matches_highlighter_table():
    document = _open(_build())
    section = document.sections[0]
    assert (section.page_width, section.page_height) == (Twips(11906), Twips(16838))
    assert section.left_margin == section.top_margin == Twips(1440)
    table = document.tables[0]
    assert table.style.name == "Table Grid"
    properties = table._tbl.tblPr
    assert properties.find(qn("w:tblLayout")).get(qn("w:type")) == "fixed"
    assert properties.find(qn("w:tblW")).get(qn("w:w")) == str(sum(COLUMN_WIDTHS))
    assert [column.w for column in table._tbl.tblGrid.gridCol_lst] == [Twips(w) for w in COLUMN_WIDTHS]


def test_docx_formatting_is_identical_to_table_highlight_output():
    """Властивості run, маркера і комірки збігаються з виводом table-highlight байт у байт."""
    ours = _open(_build(font_name="Times New Roman", font_size=12)).tables[0]
    reference = _highlighter_reference("Times New Roman", 12)
    for ours_cell, reference_cell in zip(ours.rows[0].cells, reference.rows[0].cells):
        assert _xml(ours_cell._tc.tcPr) == _xml(reference_cell._tc.tcPr)
        ours_marker, reference_marker = ours_cell.paragraphs[0], reference_cell.paragraphs[0]
        assert _xml(ours_marker._p.pPr) == _xml(reference_marker._p.pPr)
        assert _xml(ours_marker.runs[0]._r.rPr) == _xml(reference_marker.runs[0]._r.rPr)

    def by_status(table):
        found = {}
        for cell in table.rows[0].cells:
            for paragraph in cell.paragraphs[1:]:
                for run in paragraph.runs:
                    found.setdefault(_highlight(run), _xml(run._r.rPr))
        return found

    ours_runs, reference_runs = by_status(ours), by_status(reference)
    assert {"yellow", "cyan"} <= set(ours_runs)
    for status in ("yellow", "cyan"):
        assert ours_runs[status] == reference_runs[status]


def test_docx_highlights_match_yellow_and_difference_cyan_without_underline():
    table = _open(_build()).tables[0]
    runs = [run for paragraph in table.rows[0].cells[1].paragraphs[1:] for run in paragraph.runs]
    colours = {run.text.strip(): _highlight(run) for run in runs if run.text.strip()}
    assert colours["заміна"] == "cyan"
    assert "yellow" in colours.values()
    assert all(run.font.underline in (None, False) for run in runs)
    # Заливки екрана у Word немає: лише маркер виділення.
    assert next(table._tbl.iter(qn("w:shd")), None) is None


def test_docx_default_font_is_highlighter_default():
    run = _open(_build()).tables[0].rows[0].cells[0].paragraphs[1].runs[0]
    fonts = run._r.rPr.find(qn("w:rFonts"))
    assert fonts.get(qn("w:ascii")) == HighlightOptions.font_name
    assert run._r.rPr.find(qn("w:sz")).get(qn("w:val")) == str(HighlightOptions.font_size * 2)


def test_docx_keeps_raw_text_unescaped_and_long_fragment_whole_without_context():
    """HTML-екранування у Word не потрібне, довгий фрагмент не згортається, контексту немає."""
    words = [f"слово{i}" for i in range(200)]
    lines = _lines("<script> " + " ".join(words) + " & кінець")
    tokens = tokenize_lines(lines)
    segments = compare_tokens(tokens, tokens).segments
    table = _open(build_comparison_docx(
        segments, lines, tokens, lines, tokens, name_a="a", name_b="b",
    )).tables[0]
    cell_text = table.rows[0].cells[0].text
    assert "&lt;" not in cell_text
    assert "слово0" in cell_text and "слово199" in cell_text
    assert "…" not in cell_text


def test_docx_line_break_becomes_new_paragraph_and_marker_uses_first_sheet():
    lines_a = [
        {"line": " ".join(f"слово{i}" for i in range(12)), "page": 1},
        {"line": " ".join(f"слово{i}" for i in range(12, 24)), "page": 2},
    ]
    tokens_a = tokenize_lines(lines_a)
    segments = compare_tokens(tokens_a, tokens_a).segments
    pieces = fragment_pieces(
        lines_a, tokens_a, segments[0].a_start, segments[0].a_end, segments[0].a_spans, 0
    )
    # Перехід на інший аркуш PDF — розрив рядка, як <br> на екрані.
    assert ("", LINE_BREAK) in pieces
    cell = _open(build_comparison_docx(
        segments, lines_a, tokens_a, lines_a, tokens_a, name_a="a", name_b="b",
    )).tables[0].rows[0].cells[0]
    assert cell.paragraphs[0].text == "С. 1"
    assert len(cell.paragraphs) == 3
    assert next(cell._tc.iter(qn("w:br")), None) is None


def test_page_marker_without_pages_leaves_placeholder():
    tokens = tokenize_lines(_lines("одне два три"))
    assert page_marker(tokens, 0, len(tokens)) == MISSING_PAGE_MARKER


def test_docx_without_findings_says_so():
    document = _open(build_comparison_docx([], [], [], [], [], name_a="a", name_b="b"))
    assert document.tables == []
    assert EMPTY_TEXT in [paragraph.text for paragraph in document.paragraphs]
