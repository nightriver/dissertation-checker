"""Контракт тексту та оформлення звіту — PLAN_REPORT_FIRST_PAGE.md, крок A."""

from datetime import date
from html.parser import HTMLParser
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import pytest

from report_first_page.docx_export import (
    TEMPLATE_PATH, ReportTemplateError, build_report_docx, validate_template,
)
from report_first_page.presentation import (
    DATE_SPACE_BEFORE_PT, EXPERT_TEXT, HEADING_SIZE_PT, build_report_page, render_report_html,
)
from report_first_page.types import AcademicPerson, ReportMetadata, ReportPage
from tools.build_report_first_page_template import build_template

FIXTURE_PATH = Path(__file__).parent / "fixtures" / "report_first_page_expectations.json"
CORPUS = json.loads(FIXTURE_PATH.read_text(encoding="utf-8"))
EXPERT_EXPECTED = (
    "Перевірку провела Благодєтєлєва-Вовк Світлана Леонідівна, канд. екон. наук, "
    "доцент, голова Ради ГО ТРОН, координаторка антиплагіатної ініціативи «Дисергейт»"
)


def metadata_from_record(record):
    return ReportMetadata(**{
        **record,
        "supervisor": AcademicPerson(**record["supervisor"]),
        "opponents": tuple(AcademicPerson(**p) for p in record["opponents"]),
    })


class ParagraphReader(HTMLParser):
    def __init__(self):
        super().__init__()
        self.paragraphs = []
        self.bold = False
        self.active = False

    def handle_starttag(self, tag, attrs):
        if tag == "p":
            self.active = True
            self.paragraphs.append([])
        if tag == "strong":
            self.bold = True

    def handle_endtag(self, tag):
        if tag == "p":
            self.active = False
        if tag == "strong":
            self.bold = False

    def handle_data(self, data):
        if self.active:
            self.paragraphs[-1].extend((char, self.bold) for char in data)


def assert_html_word_equal(page):
    parser = ParagraphReader()
    parser.feed(render_report_html(page))
    doc = Document(BytesIO(build_report_docx(page)))
    word_fragments = [
        [(char, bool(run.bold)) for run in p.runs for char in run.text]
        for p in doc.paragraphs
    ]
    assert parser.paragraphs == word_fragments
    assert [p.text for p in doc.paragraphs] == [p.text for p in page.paragraphs]
    return doc


@pytest.mark.parametrize("filename", CORPUS)
def test_full_manual_corpus_has_identical_text_and_bold_in_html_and_word(filename):
    metadata = metadata_from_record(CORPUS[filename])
    page = build_report_page(metadata, generated_on=date(2026, 10, 21))
    doc = assert_html_word_equal(page)
    assert doc.paragraphs[2].text == (
        f"{metadata.author} на здобуття {metadata.degree_phrase} на тему «{metadata.title}» "
        "на наявність порушень академічної доброчесності"
    )
    assert len(doc.paragraphs) == 13 + len(metadata.opponents)
    assert doc.paragraphs[-2].text == EXPERT_EXPECTED
    assert doc.paragraphs[-1].text == "Київ, жовтень 2026"
    assert doc.paragraphs[-4].text == (
        f"Захист відбувся «{metadata.defense_day}» {metadata.defense_month} "
        f"{metadata.defense_year} року на засіданні спеціалізованої вченої ради "
        f"{metadata.council_code} {metadata.institution_tail}."
    )
    for paragraph, person in zip(
        (p for p in page.paragraphs if p.role in ("supervisor", "opponent")),
        (metadata.supervisor, *metadata.opponents), strict=True,
    ):
        assert [r.text for r in paragraph.runs if r.bold] == [person.full_name]
        assert person.before_name in paragraph.text
        assert person.after_name in paragraph.text


@pytest.mark.parametrize("before,after", [(None, None), ("доктор наук", None),
                                          (None, "університет"), ("доктор наук", "університет")])
def test_person_punctuation_and_bold_boundaries(before, after):
    person = AcademicPerson(before, "Коваль Олена Іванівна", after)
    metadata = ReportMetadata(supervisor=person, opponents=(person, person))
    page = build_report_page(metadata, generated_on=date(2026, 1, 1))
    doc = assert_html_word_equal(page)
    expected = (before + " " if before else "") + person.full_name + (", " + after if after else "")
    assert doc.paragraphs[5].text == expected + "."
    assert doc.paragraphs[8].text == expected + ";"
    assert doc.paragraphs[9].text == expected + "."
    assert doc.paragraphs[5].runs[1].text == person.full_name
    assert doc.paragraphs[5].runs[-1].bold is False


def test_missing_values_are_explicit_and_static_fields_do_not_change():
    page = build_report_page(ReportMetadata(), generated_on=date(2027, 1, 1))
    doc = assert_html_word_equal(page)
    assert doc.paragraphs[2].text.startswith("_____ на здобуття _____ на тему «_____»")
    assert doc.paragraphs[4].text == "Науковий керівник / консультант:"
    assert doc.paragraphs[5].text == "_____"
    assert doc.paragraphs[8].text == "_____"
    assert doc.paragraphs[10].text == (
        "Захист відбувся «_____» _____ _____ року на засіданні спеціалізованої вченої ради _____ _____."
    )
    assert doc.paragraphs[-2].text == EXPERT_EXPECTED == EXPERT_TEXT
    assert doc.paragraphs[-1].text == "Київ, січень 2027"


def test_qualification_abbreviation_keeps_its_period_before_name():
    page = build_report_page(ReportMetadata(supervisor=AcademicPerson(
        "канд. екон. наук.", "Коваль Олена Іванівна", "Університет...",
    )), generated_on=date(2026, 10, 21))
    doc = assert_html_word_equal(page)
    assert doc.paragraphs[5].text == "канд. екон. наук. Коваль Олена Іванівна, Університет."


@pytest.mark.parametrize("degree", ["наукового ступеня кандидата юридичних наук",
                                   "ступеня доктора філософії", "ступеня після ручного уточнення"])
def test_full_degree_phrase_is_used_verbatim(degree):
    page = build_report_page(ReportMetadata(degree_phrase=degree), generated_on=date(2026, 1, 1))
    doc = assert_html_word_equal(page)
    assert f"на здобуття {degree} на тему" in doc.paragraphs[2].text
    assert "ступеня ступеня" not in doc.paragraphs[2].text


@pytest.mark.parametrize("month,name", enumerate((
    "січень", "лютий", "березень", "квітень", "травень", "червень",
    "липень", "серпень", "вересень", "жовтень", "листопад", "грудень",
), 1))
def test_all_report_months_are_locale_independent(month, name):
    page = build_report_page(ReportMetadata(), generated_on=date(2028, month, 1))
    assert page.paragraphs[-1].text == f"Київ, {name} 2028"
    assert page.paragraphs[-1].runs[0].bold


def test_manual_markup_is_escaped_in_html_and_literal_in_docx():
    text = '<script>"<&\'не HTML\'</script>'
    page = build_report_page(ReportMetadata(
        author=text, title=text, degree_phrase=text, supervisor_label=text,
        supervisor=AcademicPerson(text, text, text), institution_tail=text,
    ), generated_on=date(2026, 10, 21))
    html = render_report_html(page)
    assert "<script>" not in html
    assert "&lt;script&gt;&quot;&lt;&amp;&#x27;" in html
    assert_html_word_equal(page)


def test_template_geometry_typography_and_all_slots():
    template = Document(TEMPLATE_PATH)
    validate_template(template)
    source_section_xml = template.sections[0]._sectPr.xml
    page = build_report_page(metadata_from_record(next(iter(CORPUS.values()))), generated_on=date(2026, 10, 21))
    doc = assert_html_word_equal(page)
    assert doc.sections[0]._sectPr.xml == source_section_xml
    section = doc.sections[0]
    assert (section.page_width.twips, section.page_height.twips) == (11906, 16838)
    assert all(m.twips == 1440 for m in (section.top_margin, section.left_margin, section.right_margin, section.bottom_margin))
    assert doc.paragraphs[-1].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert len(doc.paragraphs[-1].runs) == 1 and doc.paragraphs[-1].runs[0].bold
    assert doc.paragraphs[-1].paragraph_format.space_before.pt == DATE_SPACE_BEFORE_PT == 17
    assert f"margin-top:{DATE_SPACE_BEFORE_PT}pt" in render_report_html(page)
    for index, p in enumerate(doc.paragraphs):
        assert p.paragraph_format.space_before.pt == (DATE_SPACE_BEFORE_PT if index == len(doc.paragraphs) - 1 else 0)
        assert p.paragraph_format.space_after.pt == 0
        assert p.paragraph_format.line_spacing == 1
        for r in p.runs:
            assert r.font.name == ("Calibri Light" if index < 3 else "Calibri")
            assert r.font.size.pt == (HEADING_SIZE_PT if index < 3 else 14)
            assert str(r.font.color.rgb) == "000000"
    xml = doc.element.xml + doc.styles.element.xml
    assert "TimesNewRomanPS" not in xml
    assert "{{" not in doc.element.xml and "}}" not in doc.element.xml
    with ZipFile(BytesIO(build_report_docx(page))) as package:
        assert all(b"TimesNewRomanPS" not in package.read(name)
                   for name in package.namelist() if name.endswith(".xml"))


@pytest.mark.parametrize("damage", ["count", "split", "duplicate", "unknown"])
def test_template_validation_rejects_word_split_and_unknown_slots(tmp_path, damage):
    doc = Document(TEMPLATE_PATH)
    if damage == "count":
        doc.add_paragraph()
    elif damage == "split":
        run = doc.paragraphs[2].runs[0]
        run.text = "{{AUTH"
        doc.paragraphs[2].add_run("OR}}")
    elif damage == "duplicate":
        doc.paragraphs[2].add_run("{{AUTHOR}}")
    else:
        doc.paragraphs[2].add_run("{{UNEXPECTED}}")
    path = tmp_path / "broken.docx"
    doc.save(path)
    with pytest.raises(ReportTemplateError):
        build_report_docx(build_report_page(ReportMetadata(), generated_on=date(2026, 1, 1)), template_path=path)


def test_missing_template_and_invalid_page_fail_explicitly(tmp_path):
    page = build_report_page(ReportMetadata(), generated_on=date(2026, 1, 1))
    with pytest.raises(ReportTemplateError, match="відкрити"):
        build_report_docx(page, template_path=tmp_path / "absent.docx")
    with pytest.raises(ValueError, match="структура"):
        build_report_docx(ReportPage(()))


def test_template_builder_preserves_reference_and_theme(tmp_path):
    source = tmp_path / "source.docx"
    document = Document(TEMPLATE_PATH)
    document.core_properties.author = "Автор іншої дисертації"
    document.save(source)
    original = source.read_bytes()
    output = tmp_path / "prepared.docx"
    build_template(source, output)
    assert source.read_bytes() == original
    validate_template(Document(output))
    assert not Document(output).core_properties.author
    with ZipFile(source) as before, ZipFile(output) as after:
        assert before.read("word/theme/theme1.xml") == after.read("word/theme/theme1.xml")
    with pytest.raises(ValueError, match="перезаписувати"):
        build_template(source, source)
    document.paragraphs[12].text = "Не той експерт"
    document.save(source)
    with pytest.raises(ValueError, match="експерта"):
        build_template(source, output)


def test_export_does_not_depend_on_examples_or_working_directory(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    page = build_report_page(ReportMetadata(), generated_on=date(2026, 10, 21))
    assert Document(BytesIO(build_report_docx(page))).paragraphs[-1].text == "Київ, жовтень 2026"
