"""Стан, правки й маршрут екрана — PLAN_REPORT_FIRST_PAGE.md, крок C."""

from dataclasses import replace
from datetime import date, datetime, timezone
from io import BytesIO
import json
from pathlib import Path

from docx import Document
import pytest
from streamlit.testing.v1 import AppTest

from report_first_page import ui
from report_first_page.extractor import ReportInputError
from report_first_page.presentation import AUTHOR_WARNING
from report_first_page.types import AcademicPerson, ReportMetadata
from ui_helpers import is_report_first_page_mode

ROOT = Path(__file__).resolve().parents[1]
APP_PATH = ROOT / "app.py"
CORPUS = json.loads((ROOT / "tests/fixtures/report_first_page_expectations.json").read_text(encoding="utf-8"))


def sample_metadata():
    record = next(iter(CORPUS.values()))
    return ReportMetadata(**{
        **record, "supervisor": AcademicPerson(**record["supervisor"]),
        "opponents": tuple(AcademicPerson(**p) for p in record["opponents"]),
    })


def open_screen():
    app = AppTest.from_file(APP_PATH)
    app.query_params["mode"] = "report-first-page"
    app.run(timeout=30)
    assert not app.exception
    return app


def preview(app):
    return next(item.value for item in app.markdown if '<article class="report-first-page">' in item.value)


@pytest.mark.parametrize("value,expected", [("report-first-page", True), (["report-first-page"], True),
    (("report-first-page",), True), ([], False), (None, False), ("search", False)])
def test_report_query_parameter_supports_scalar_and_list(value, expected):
    assert is_report_first_page_mode({"mode": value}) is expected
    assert is_report_first_page_mode({}) is False


def test_report_screen_is_last_and_has_exactly_one_pdf_uploader():
    app = open_screen()
    assert [t.value for t in app.title] == ["Звіт"]
    assert [item.label for item in app.get("file_uploader")] == ["Автореферат (PDF із текстовим шаром)"]
    assert not app.get("download_button")
    navigation = next(m.value for m in app.markdown if 'class="app-main-nav"' in m.value)
    assert 'app-main-nav__item--active" href="?mode=report-first-page" target="_self" aria-current="page">Звіт' in navigation
    assert navigation.endswith('>Звіт</a></nav>')


def test_upload_preview_download_edits_and_date_reruns_do_not_reparse(monkeypatch):
    calls = []
    downloads = []
    toasts = []
    generated = [date(2026, 10, 31)]
    metadata = sample_metadata()
    monkeypatch.setattr(ui, "extract_report_metadata", lambda data: calls.append(data) or metadata)
    monkeypatch.setattr(ui, "kyiv_date", lambda: generated[0])
    original_toast = ui.st.toast
    def toast(body, **kwargs):
        toasts.append((body, kwargs))
        return original_toast(body, **kwargs)
    monkeypatch.setattr(ui.st, "toast", toast)
    original_docx = ui.build_report_docx
    def export(page):
        data = original_docx(page)
        downloads.append(data)
        return data
    monkeypatch.setattr(ui, "build_report_docx", export)
    app = open_screen()
    app.get("file_uploader")[0].upload("same.pdf", b"PDF A", "application/pdf")
    app.run(timeout=30)
    assert not app.exception
    assert toasts == [("Перевірте відмінок ПІБ автора", {"icon": "ℹ️", "duration": 5})]
    assert AUTHOR_WARNING not in [w.value for w in app.warning]
    assert len(app.get("download_button")) == 1
    assert app.get("download_button")[0].label == "Завантажити"
    assert app.get("download_button")[0].proto.type == "primary"
    assert app.expander[0].label == "Уточнити дані" and app.expander[0].proto.expanded is False
    assert not app.get("image") and not app.get("slider")
    assert preview(app).count(metadata.author) == 1
    left, right = app.get("column")
    assert preview(app) in [item.value for item in left.markdown]
    assert not left.get("download_button") and not left.get("expander")
    assert len(right.get("download_button")) == len(right.get("expander")) == 1
    ordered = list(right)
    download_index = next(i for i, element in enumerate(ordered) if element.type == "download_button")
    edit_index = next(i for i, element in enumerate(ordered) if element.type == "expander")
    assert download_index < edit_index
    app.text_input(key=ui.FIELD_PREFIX + "author").set_value("Автор після правки").run(timeout=30)
    assert "Автор після правки на здобуття" in preview(app)
    assert "Автор після правки на здобуття" in Document(BytesIO(downloads[-1])).paragraphs[2].text
    app.text_area(key=ui.FIELD_PREFIX + "title").set_value("").run(timeout=30)
    assert "на тему «_____»" in preview(app)
    app.text_input(key=ui.FIELD_PREFIX + "opponents_0_full_name").set_value("Перший після правки").run(timeout=30)
    assert "Перший після правки" in preview(app)
    assert app.text_input(key=ui.FIELD_PREFIX + "opponents_1_full_name").value == metadata.opponents[1].full_name
    previous = downloads[-1]
    generated[0] = date(2026, 11, 1)
    app.run(timeout=30)
    assert "Київ, листопад 2026" in preview(app)
    assert previous != downloads[-1]
    assert Document(BytesIO(downloads[-1])).paragraphs[-1].text == "Київ, листопад 2026"
    assert len(calls) == 1
    assert len(toasts) == 1


def test_every_variable_field_updates_both_outputs_immediately(monkeypatch):
    captured = []
    monkeypatch.setattr(ui, "extract_report_metadata", lambda _: sample_metadata())
    original = ui.build_report_docx
    def export(page):
        captured.append(page)
        return original(page)
    monkeypatch.setattr(ui, "build_report_docx", export)
    app = open_screen()
    app.get("file_uploader")[0].upload("source.pdf", b"data", "application/pdf")
    app.run(timeout=30)
    assert {w.key.removeprefix(ui.FIELD_PREFIX) for w in app.text_area} == {
        "title", "degree_phrase", "institution_tail", "supervisor_before_name", "supervisor_after_name",
        "opponents_0_before_name", "opponents_0_after_name", "opponents_1_before_name", "opponents_1_after_name",
    }
    for index, widget in enumerate([*app.text_input, *app.text_area]):
        multiline = widget.type == "text_area"
        text = f"Поле {index}\nпісля правки" if multiline else f"Поле {index} після правки"
        collection = app.text_area if multiline else app.text_input
        collection(key=widget.key).set_value(text).run(timeout=30)
        assert not app.exception
        assert text in preview(app)
        assert any(text in p.text for p in captured[-1].paragraphs)


def test_same_name_and_size_file_switch_invalid_file_removal_and_return_reset_edits(monkeypatch):
    calls = []
    def extract(data):
        calls.append(data)
        if data == b"ERROR":
            raise ReportInputError("Пошкоджений PDF B")
        return replace(sample_metadata(), author="Автор А" if data == b"PDF A" else "Автор Б")
    monkeypatch.setattr(ui, "extract_report_metadata", extract)
    app = open_screen()
    uploader = app.get("file_uploader")[0]
    uploader.upload("same.pdf", b"PDF A", "application/pdf")
    app.run(timeout=30)
    app.text_input(key=ui.FIELD_PREFIX + "author").set_value("Ручна правка А").run(timeout=30)
    app.get("file_uploader")[0].upload("same.pdf", b"PDF B", "application/pdf")
    app.run(timeout=30)
    assert "Автор Б" in preview(app) and "Ручна правка А" not in preview(app)
    app.get("file_uploader")[0].upload("same.pdf", b"ERROR", "application/pdf")
    app.run(timeout=30)
    assert "Пошкоджений PDF B" in [e.value for e in app.error]
    assert not app.text_input and not app.text_area and not app.get("download_button")
    assert not any("report-first-page\">" in m.value for m in app.markdown)
    app.run(timeout=30)
    assert calls.count(b"ERROR") == 1
    app.get("file_uploader")[0].upload("same.pdf", b"PDF A", "application/pdf")
    app.run(timeout=30)
    assert app.text_input(key=ui.FIELD_PREFIX + "author").value == "Автор А"
    assert ui.AUTHOR_TOAST_KEY in app.session_state
    app.get("file_uploader")[0].clear()
    app.run(timeout=30)
    assert not app.get("download_button") and not app.text_input and not app.text_area
    assert ui.METADATA_KEY not in app.session_state
    assert ui.AUTHOR_TOAST_KEY not in app.session_state


def test_docx_export_error_keeps_html_and_editing_but_removes_download(monkeypatch):
    monkeypatch.setattr(ui, "extract_report_metadata", lambda _: sample_metadata())
    def fail(page):
        raise ValueError("Відсутній шаблон")
    monkeypatch.setattr(ui, "build_report_docx", fail)
    app = open_screen()
    app.get("file_uploader")[0].upload("source.pdf", b"data", "application/pdf")
    app.run(timeout=30)
    assert not app.exception
    assert preview(app) and app.text_input and app.text_area
    assert not app.get("download_button")
    assert any("Відсутній шаблон" in e.value for e in app.error)


def test_invalid_upload_clears_existing_result_before_parsing(monkeypatch):
    monkeypatch.setattr(ui, "extract_report_metadata", lambda _: sample_metadata())
    app = open_screen()
    app.get("file_uploader")[0].upload("source.pdf", b"data", "application/pdf")
    app.run(timeout=30)
    assert app.get("download_button")
    app.get("file_uploader")[0].upload("source.pdf", b"", "application/pdf")
    app.run(timeout=30)
    assert not app.exception and app.error
    assert not app.text_input and not app.text_area and not app.get("download_button")
    assert ui.METADATA_KEY not in app.session_state


def test_missing_fields_remain_empty_in_editor_and_can_be_filled(monkeypatch):
    monkeypatch.setattr(ui, "extract_report_metadata", lambda _: ReportMetadata(warnings=("Немає даних",)))
    app = open_screen()
    app.get("file_uploader")[0].upload("source.pdf", b"data", "application/pdf")
    app.run(timeout=30)
    assert not app.exception
    assert all(w.value == "" for w in [*app.text_input, *app.text_area])
    assert "Немає даних" in [w.value for w in app.warning]
    app.text_input(key=ui.FIELD_PREFIX + "opponents_0_full_name").set_value("Заповнено вручну").run(timeout=30)
    assert "Заповнено вручну" in preview(app)


@pytest.mark.corpus
def test_real_pdf_automatically_renders_complete_report_without_confirmation(monkeypatch):
    monkeypatch.setattr(ui, "kyiv_date", lambda: date(2026, 10, 21))
    app = open_screen()
    data = (ROOT / "examples/avtoreferats/autoreferat_Автореферат_Бенедесюк_І_М.pdf").read_bytes()
    app.get("file_uploader")[0].upload("wrong-spelling.pdf", data, "application/pdf")
    app.run(timeout=30)
    assert not app.exception
    html = preview(app)
    assert "Бенедисюк Ігор Михайлович" in html
    assert "«13» травня 2021 року" in html
    assert "К 26.504.01 у Науково-дослідному інституті" in html
    assert "Гетманцев Максим Олександрович" in html
    assert "Украї–ни" not in html
    assert not app.button and app.get("download_button")


@pytest.mark.parametrize("utc,expected", [
    (datetime(2026, 10, 31, 22, 30, tzinfo=timezone.utc), date(2026, 11, 1)),
    (datetime(2026, 12, 31, 22, 30, tzinfo=timezone.utc), date(2027, 1, 1)),
    (datetime(2026, 6, 30, 21, 30, tzinfo=timezone.utc), date(2026, 7, 1)),
])
def test_kyiv_date_uses_timezone_at_month_year_and_dst_boundaries(utc, expected):
    assert ui.kyiv_date(clock=lambda tz: utc) == expected


def test_source_reset_keeps_uploader_and_unrelated_modes_only():
    state = {ui.UPLOAD_KEY: "upload", ui.METADATA_KEY: "old", ui.FIELD_PREFIX + "author": "edited",
             ui.AUTHOR_TOAST_KEY: True, "search_result": "keep"}
    ui.reset_source(state, "new")
    assert state == {ui.UPLOAD_KEY: "upload", ui.SOURCE_KEY: "new", "search_result": "keep"}
    ui.reset_source(state, None)
    assert state == {ui.UPLOAD_KEY: "upload", "search_result": "keep"}
