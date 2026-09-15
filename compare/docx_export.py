"""
Вивантаження таблиці текстових збігів у документ MS Word (.docx).

Таблиця оформлюється так само, як результат режиму ``?mode=table-highlight``:
експерт збирає обидві таблиці в одному висновку, тому шрифт, геометрія,
маркер «С. N» і кольори підсвічування мають збігатися. Оформлення не
дублюється, а береться з функцій ``table_highlighter``.
"""

from __future__ import annotations

from collections.abc import Sequence
from io import BytesIO

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Twips

from compare.presentation import (
    LINE_BREAK,
    SIDE_A_TITLE,
    SIDE_B_TITLE,
    fragment_pieces,
)
from compare.types import CompareToken, TextSegment
from parser.types import LineItem
from table_highlighter.formatting import normalize_document, set_run_font
from table_highlighter.layout import LogicalRow, align_page_markers
from table_highlighter.types import HighlightOptions
from table_highlighter.writer import HIGHLIGHT
from table_highlighter.zones import CellZones, ParagraphZone


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
DOCUMENT_TITLE = "Порівняння двох робіт: текстові збіги"
EMPTY_TEXT = "Знахідок немає."
MISSING_PAGE_MARKER = "С. —"
# Геометрія таблиці порівняння, яку експерт обробляє в table-highlight:
# A4 книжкова, поля 2,54 см, дві колонки фіксованої ширини (twips).
PAGE_WIDTH = 11906
PAGE_HEIGHT = 16838
PAGE_MARGIN = 1440
COLUMN_WIDTHS = (4673, 4678)
# Інтервали абзаців поза таблицею — як у стандартному документі Word.
SPACING_AFTER = 160
LINE_SPACING = 259

_STATUS = {"equal": "match", "fuzzy": "match", "replace": "diff", "insert": "diff", "delete": "diff"}


def page_marker(tokens: Sequence[CompareToken], start: int, end: int) -> str:
    """Лише перший аркуш PDF фрагмента; назву документа експерт допише сам."""
    for token in tokens[start:end]:
        for page in token.physical_pages:
            if page is not None:
                return f"С. {page}"
    return MISSING_PAGE_MARKER


def _set_paragraph_defaults(document) -> None:
    defaults = document.styles.element.find(qn("w:docDefaults"))
    spacing = defaults.find(qn("w:pPrDefault")).find(qn("w:pPr")).find(qn("w:spacing"))
    spacing.set(qn("w:after"), str(SPACING_AFTER))
    spacing.set(qn("w:line"), str(LINE_SPACING))
    spacing.set(qn("w:lineRule"), "auto")


def _prepare_table(document):
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    properties = table._tbl.tblPr
    width = properties.find(qn("w:tblW"))
    width.set(qn("w:w"), str(sum(COLUMN_WIDTHS)))
    width.set(qn("w:type"), "dxa")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    width.addnext(layout)
    for column, value in zip(table._tbl.tblGrid.gridCol_lst, COLUMN_WIDTHS):
        column.w = Twips(value)
    return table


def _merge_same_status(groups: list[list]) -> list[list]:
    merged: list[list] = []
    for text, status in groups:
        if merged and merged[-1][1] == status:
            merged[-1][0] += text
        else:
            merged.append([text, status])
    return merged


def _paragraph_runs(pieces) -> list[list[list]]:
    """
    Розкладає фрагмент на абзаци з мінімальною кількістю runs.

    Сусідні шматки одного статусу стають одним run: кожен run у Word несе
    повний набір властивостей шрифту, і пословні runs для фрагмента на
    десятки тисяч слів займали гігабайти пам'яті.

    Шматок без літер і цифр (пробіли, розділові знаки, лапки, дужки, тире)
    бере статус сусідів, якщо вони однакові, а на краю абзацу — статус
    єдиного сусіда. Розриви підсвічування на «», » (» лише створюють
    візуальний шум. Між збігом і відмінністю такий шматок лишається без кольору.
    """
    paragraphs: list[list[list]] = [[]]
    for text, operation in pieces:
        if operation == LINE_BREAK:
            paragraphs.append([])
        elif text:
            paragraphs[-1].append([text, _STATUS.get(operation)])
    result = []
    for groups in paragraphs:
        groups = _merge_same_status(groups)
        for index, (text, status) in enumerate(groups):
            if status is not None or any(char.isalnum() for char in text):
                continue
            neighbours = {
                groups[position][1]
                for position in (index - 1, index + 1)
                if 0 <= position < len(groups)
            }
            if len(neighbours) == 1 and None not in neighbours:
                groups[index][1] = neighbours.pop()
        result.append(_merge_same_status(groups))
    return result


def _fill_cell(cell, marker: str, pieces, highlighted: list) -> CellZones:
    """Абзац «С. N», під ним фрагмент; розрив рядка стає новим абзацом."""
    cell.paragraphs[0].add_run(marker)
    for groups in _paragraph_runs(pieces):
        paragraph = cell.add_paragraph()
        for text, status in groups:
            run = paragraph.add_run(text)
            if status is not None:
                highlighted.append((run._r, status))
    zones = tuple(
        ParagraphZone(item, index, "plain" if index == 0 else "text")
        for index, item in enumerate(cell.paragraphs)
    )
    return CellZones(zones, marker_index=0)


def build_comparison_docx(
    segments: Sequence[TextSegment],
    lines_a: Sequence[LineItem], tokens_a: Sequence[CompareToken],
    lines_b: Sequence[LineItem], tokens_b: Sequence[CompareToken],
    *,
    name_a: str,
    name_b: str,
    summary: Sequence[str] = (),
    font_name: str = HighlightOptions.font_name,
    font_size: int = HighlightOptions.font_size,
) -> bytes:
    """
    Будує .docx з тими самими знахідками, що й таблиця на екрані.

    Кожна знахідка — один рядок таблиці з двох комірок. У комірці лише
    маркер «С. N» і сам фрагмент без контексту, повністю, без згортання.
    Збіг (точний і близька словоформа) підсвічено жовтим маркером Word,
    відмінність — бірюзовим, як у table-highlight.
    """
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Twips(PAGE_WIDTH), Twips(PAGE_HEIGHT)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Twips(PAGE_MARGIN))
    _set_paragraph_defaults(document)

    document.add_paragraph(DOCUMENT_TITLE)
    document.add_paragraph(f"{SIDE_A_TITLE}: {name_a}")
    document.add_paragraph(f"{SIDE_B_TITLE}: {name_b}")
    for line in summary:
        document.add_paragraph(line)

    if not segments:
        document.add_paragraph(EMPTY_TEXT)
    highlighted: list = []
    rows = []
    table = _prepare_table(document) if segments else None
    for segment in segments:
        row = table.add_row()
        for cell, value in zip(row.cells, COLUMN_WIDTHS):
            cell.width = Twips(value)
        left = _fill_cell(
            row.cells[0], page_marker(tokens_a, segment.a_start, segment.a_end),
            fragment_pieces(lines_a, tokens_a, segment.a_start, segment.a_end, segment.a_spans, 0),
            highlighted,
        )
        right = _fill_cell(
            row.cells[1], page_marker(tokens_b, segment.b_start, segment.b_end),
            fragment_pieces(lines_b, tokens_b, segment.b_start, segment.b_end, segment.b_spans, 0),
            highlighted,
        )
        rows.append((row, left, right))

    # Той самий порядок, що в table_highlighter.processor: спершу єдиний
    # шрифт на весь документ, потім підсвічування, потім вирівнювання маркерів.
    normalize_document(document, font_name, font_size)
    for run, status in highlighted:
        set_run_font(run, font_name, font_size, HIGHLIGHT[status])
    for row, left, right in rows:
        align_page_markers(LogicalRow(row), left, right, font_name, font_size)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
