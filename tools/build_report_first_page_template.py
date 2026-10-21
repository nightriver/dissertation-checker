"""Відтворювана підготовка шаблону — PLAN_REPORT_FIRST_PAGE.md, §3."""

import argparse
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from lxml import etree

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from report_first_page.docx_export import SLOTS, format_run, validate_template
from report_first_page.presentation import DATE_SPACE_BEFORE_PT, EXPERT_TEXT


def build_template(source: Path, output: Path) -> None:
    """Зберігає тему й секцію зразка; тіло нормалізує за контрактом."""
    if source.resolve() == output.resolve():
        raise ValueError("Зразок не можна перезаписувати.")
    document = Document(source)
    if not any(p.text == EXPERT_TEXT for p in document.paragraphs):
        raise ValueError("У зразку немає погодженого повного абзацу експерта.")
    body = document.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    # Прибираємо PostScript-назви також зі стилів, а не лише з видимого тексту.
    for fonts in document.styles.element.iter(qn("w:rFonts")):
        for attr, value in list(fonts.attrib.items()):
            if "TimesNewRomanPS" in value:
                fonts.set(attr, "Calibri")
    # Старі записи більше ніде не використовуються; Calibri вже є в таблиці.
    font_part = document.part.part_related_by(RELATIONSHIP_TYPE.FONT_TABLE)
    font_table = etree.fromstring(font_part.blob)
    for font in list(font_table):
        if "TimesNewRomanPS" in font.get(qn("w:name"), ""):
            font_table.remove(font)
    font_part._blob = etree.tostring(font_table, xml_declaration=True, encoding="UTF-8", standalone=True)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(14)
    normal.font.bold = False
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    static = {0: "ЗВІТ", 1: "результатів перевірки дисертації", 7: "Опоненти:"}
    slot_texts = {
        2: ("{{AUTHOR}}", " на здобуття ", "{{DEGREE_PHRASE}}", " на тему «", "{{TITLE}}",
            "» на наявність порушень академічної доброчесності"),
        4: ("{{SUPERVISOR_LABEL}}", ":"),
        5: ("{{SUPERVISOR_BEFORE}}", "{{SUPERVISOR_NAME}}", "{{SUPERVISOR_AFTER}}", "{{SUPERVISOR_END}}"),
        8: ("{{OPPONENT_BEFORE}}", "{{OPPONENT_NAME}}", "{{OPPONENT_AFTER}}", "{{OPPONENT_END}}"),
        10: ("Захист відбувся", " ", "{{DEFENSE_DATE}}",
             " на засіданні спеціалізованої вченої ради ", "{{COUNCIL_CODE}}", " ", "{{INSTITUTION_TAIL}}", "."),
        13: ("{{REPORT_DATE}}",),
    }
    for index in range(14):
        paragraph = document.add_paragraph(style="Normal")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 13 else WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.line_spacing = 1
        paragraph.paragraph_format.space_before = Pt(DATE_SPACE_BEFORE_PT if index == 13 else 0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.keep_with_next = False
        paragraph.paragraph_format.keep_together = False
        paragraph.paragraph_format.widow_control = True
        if index in SLOTS:
            for text in slot_texts[index]:
                format_run(paragraph.add_run(text), heading=index < 3,
                           bold=text.endswith("_NAME}}") or index in (4, 13) or text == "Захист відбувся")
        elif index == 12:
            format_run(paragraph.add_run("Перевірку провела"), bold=True)
            format_run(paragraph.add_run(EXPERT_TEXT[len("Перевірку провела"):]))
        else:
            format_run(paragraph.add_run(static.get(index, "")), heading=index < 3, bold=index == 7)
    properties = document.core_properties
    properties.author = ""
    properties.last_modified_by = ""
    properties.title = "Звіт"
    properties.subject = ""
    properties.comments = ""
    properties.keywords = ""
    validate_template(document)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build_template(args.source, args.output)
