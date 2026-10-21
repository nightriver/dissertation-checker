"""Шаблонний Word зі спільного тексту — PLAN_REPORT_FIRST_PAGE.md, §§3, 6."""

from copy import deepcopy
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from .presentation import DATE_SPACE_BEFORE_PT, HEADING_SIZE_PT
from .types import ReportPage

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TEMPLATE_PATH = Path(__file__).resolve().parent / "resources" / "report_first_page.docx"
SLOTS = {
    2: ("AUTHOR", "DEGREE_PHRASE", "TITLE"),
    4: ("SUPERVISOR_LABEL",),
    5: ("SUPERVISOR_BEFORE", "SUPERVISOR_NAME", "SUPERVISOR_AFTER", "SUPERVISOR_END"),
    8: ("OPPONENT_BEFORE", "OPPONENT_NAME", "OPPONENT_AFTER", "OPPONENT_END"),
    10: ("DEFENSE_DATE", "COUNCIL_CODE", "INSTITUTION_TAIL"),
    13: ("REPORT_DATE",),
}
ROLES = (
    "heading", "heading", "heading", "separator", "label", "supervisor", "separator",
    "label", "opponent", "separator", "defense", "separator", "expert", "date",
)


class ReportTemplateError(ValueError):
    """Робочий шаблон відсутній або має пошкоджені слоти."""


def format_run(run, *, heading: bool = False, bold: bool = False) -> None:
    run.font.name = "Calibri Light" if heading else "Calibri"
    run.font.size = Pt(HEADING_SIZE_PT if heading else 14)
    run.font.bold = bold
    run.font.italic = False
    run.font.color.rgb = RGBColor(0, 0, 0)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + attr), run.font.name)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme", "csTheme"):
        fonts.attrib.pop(qn("w:" + attr), None)


def validate_template(document) -> None:
    """Кожен слот — рівно один run з одним w:t у своєму абзаці."""
    if len(document.paragraphs) != len(ROLES):
        raise ReportTemplateError("Шаблон звіту має містити 14 канонічних абзаців.")
    all_text = "".join(node.text or "" for node in document.element.iter(qn("w:t")))
    known = {"{{" + slot + "}}" for slots in SLOTS.values() for slot in slots}
    for index, slots in SLOTS.items():
        paragraph = document.paragraphs[index]
        for slot in slots:
            marker = "{{" + slot + "}}"
            matches = [r for r in paragraph.runs if r.text == marker]
            if (all_text.count(marker) != 1 or len(matches) != 1
                    or len(matches[0]._element.findall(qn("w:t"))) != 1):
                raise ReportTemplateError(f"Пошкоджений слот шаблону: {marker}")
    for marker in known:
        all_text = all_text.replace(marker, "")
    if "{{" in all_text or "}}" in all_text:
        raise ReportTemplateError("Невідомий слот шаблону звіту.")


def build_report_docx(page: ReportPage, *, template_path: Path | None = None) -> bytes:
    """Клонує прототип опонента; фрази та виділення бере лише з ReportPage."""
    try:
        document = Document(template_path or TEMPLATE_PATH)
    except Exception as exc:
        raise ReportTemplateError("Не вдалося відкрити робочий шаблон звіту.") from exc
    validate_template(document)
    opponent_count = sum(p.role == "opponent" for p in page.paragraphs)
    expected_roles = ROLES[:8] + ("opponent",) * opponent_count + ROLES[9:]
    if opponent_count < 1 or tuple(p.role for p in page.paragraphs) != expected_roles:
        raise ValueError("Некоректна структура готової сторінки звіту.")
    prototype = document.paragraphs[8]._p
    cursor = prototype
    for _ in range(opponent_count - 1):
        clone = deepcopy(prototype)
        cursor.addnext(clone)
        cursor = clone
    for target, source in zip(document.paragraphs, page.paragraphs, strict=True):
        target.clear()
        if source.role == "date":
            target.paragraph_format.space_before = Pt(DATE_SPACE_BEFORE_PT)
        for fragment in source.runs:
            run = target.add_run(fragment.text)
            format_run(run, heading=source.role == "heading", bold=fragment.bold)
    output = BytesIO()
    document.save(output)
    return output.getvalue()
